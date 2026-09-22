import os
import tempfile
import unittest
from openpyxl import Workbook, load_workbook

from modules.date_comparison import autofill
from utils.column_mapper import ColumnMapper

HEADERS = ColumnMapper.OUTPUT_COLUMNS

DRIVE_URL = "https://drive.google.com/file/d/Abc123XYZ/view?usp=sharing"
PROJECT = "FORTALECIMIENTO AGRICOLA"
PROJECT_SRC = "Fortalecimiento agrícola"


class AutofillTests(unittest.TestCase):

    def _build_wb(self):
        wb = Workbook()
        wb.remove(wb.active)
        return wb

    def _add_sheet(self, wb, name, rows):
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        for row in rows:
            ws.append([row.get(h) for h in HEADERS])
        return ws

    def _source_row(self):
        return {
            "Facultad": "CIENCIAS AGRÍCOLAS",
            "Carrera": "AGRONOMÍA",
            "TipoProyecto": "VIGENTE",
            "NombreProyecto": PROJECT_SRC,
            "NombrePrograma": "Fortalecimiento Integral",
            "NombreCoordinador": "Ana López",
            "Territorio": "Quito",
            "FechaInicio": "01/06/2024",
            "FechaFin": "01/06/2025",
            "LinkPlanificacion": DRIVE_URL,
            "idCodigo": "P3AGR18",
        }

    def _target_row(self, project_type):
        return {
            "Facultad": "CIENCIAS AGRICOLAS",
            "Carrera": "AGRONOMÍA",
            "TipoProyecto": project_type,
            "NombreProyecto": PROJECT,
            "NombrePrograma": "N/A",
            "NombreCoordinador": "",
            "Territorio": "Quito",
            "FechaInicio": "",
            "FechaFin": "N/A",
            "LinkPlanificacion": "N/A",
            "idCodigo": "",
        }

    def _project_index(self, wb):
        headers = ColumnMapper.OUTPUT_COLUMNS
        return headers.index("NombreProyecto")

    def _rows(self, ws):
        headers = ColumnMapper.OUTPUT_COLUMNS
        return [
            dict(zip(headers, ["" if cell.value is None else cell.value for cell in row]))
            for row in ws.iter_rows(min_row=2)
        ]

    def test_vigente_filled_from_other_sheet(self):
        wb = self._build_wb()
        self._add_sheet(wb, "2025-2026", [self._source_row()])
        self._add_sheet(wb, "2026-2026", [self._target_row("VIGENTE")])

        with tempfile.TemporaryDirectory() as tmp:
            wb_path = os.path.join(tmp, "BDVinculacionn.xlsx")
            rep_path = os.path.join(tmp, "date_comparison_report.md")
            wb.save(wb_path)
            filled = autofill.run(wb_path, rep_path)

            out = load_workbook(wb_path)
            target = self._rows(out["2026-2026"])[0]
            out.close()
            self.assertGreater(filled, 0)
            self.assertEqual(target["idCodigo"], "P3AGR18")
            self.assertEqual(target["LinkPlanificacion"], DRIVE_URL)
            self.assertEqual(target["NombrePrograma"], "Fortalecimiento Integral")
            self.assertEqual(target["NombreCoordinador"], "Ana López")
            self.assertEqual(target["FechaInicio"], "01/06/2024")
            self.assertEqual(target["FechaFin"], "01/06/2025")

            filled_wb = load_workbook(wb_path)
            target_ws = filled_wb["2026-2026"]
            link_cell = target_ws.cell(2, HEADERS.index("LinkPlanificacion") + 1)
            self.assertEqual(link_cell.hyperlink.target, DRIVE_URL)
            filled_wb.close()

            with open(rep_path, encoding="utf-8") as f:
                self.assertIn("## Autofill", f.read())

    def test_nuevo_skipped(self):
        wb = self._build_wb()
        self._add_sheet(wb, "2025-2026", [self._source_row()])
        self._add_sheet(wb, "2026-2026", [self._target_row("NUEVO")])

        with tempfile.TemporaryDirectory() as tmp:
            wb_path = os.path.join(tmp, "BDVinculacionn.xlsx")
            wb.save(wb_path)
            filled = autofill.run(wb_path, os.path.join(tmp, "report.md"))

            out = load_workbook(wb_path)
            target = self._rows(out["2026-2026"])[0]
            out.close()
            self.assertEqual(filled, 0)
            self.assertEqual(target["idCodigo"], "")
            self.assertEqual(target["LinkPlanificacion"], "N/A")

    def test_incomplete_source_date_not_filled(self):
        src = self._source_row()
        src["FechaInicio"] = "Abril 2025"
        wb = self._build_wb()
        self._add_sheet(wb, "2025-2026", [src])
        self._add_sheet(wb, "2026-2026", [self._target_row("VIGENTE")])

        with tempfile.TemporaryDirectory() as tmp:
            wb_path = os.path.join(tmp, "BDVinculacionn.xlsx")
            wb.save(wb_path)
            filled = autofill.run(wb_path, os.path.join(tmp, "report.md"))

            out = load_workbook(wb_path)
            target = self._rows(out["2026-2026"])[0]
            out.close()
            self.assertEqual(target["FechaInicio"], "")
            self.assertGreaterEqual(filled, 1)

    def test_same_sheet_row_not_used_as_own_source(self):
        row = self._source_row()
        wb = self._build_wb()
        self._add_sheet(wb, "2026-2026", [row])
        with tempfile.TemporaryDirectory() as tmp:
            wb_path = os.path.join(tmp, "BDVinculacionn.xlsx")
            wb.save(wb_path)
            filled = autofill.run(wb_path, os.path.join(tmp, "report.md"))
            self.assertEqual(filled, 0)


if __name__ == "__main__":
    unittest.main()