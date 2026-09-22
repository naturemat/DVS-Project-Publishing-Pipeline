import os
import re
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from utils.column_mapper import ColumnMapper
from utils.data_formatter import DataFormatter
from utils.facultad_matcher import FacultadMatcher
from utils.file_reader import FileExtractor

import extract_data

_FULL_DATE = re.compile(r"^\d{2}/\d{2}/\d{4}$")


class PipelineIntegrationTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        if not os.path.exists(extract_data.BD_SOURCE):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "2024-2025"
            ws.append(ColumnMapper.OUTPUT_COLUMNS)
            wb.save(extract_data.BD_SOURCE)
            wb.close()

    def _examples(self):
        extractor = FileExtractor(extract_data.BASE_DIR)
        return [f for f in extractor.list_files() if f.startswith("EJEMPLO")]

    def test_example_files_found(self):
        self.assertGreaterEqual(len(self._examples()), 3)

    def test_example_rows_are_formatted(self):
        programas = extract_data.load_programas()
        matcher = FacultadMatcher()
        for filename in self._examples():
            with self.subTest(filename=filename):
                period = ColumnMapper.detect_period(filename)
                col_map = ColumnMapper.get_column_map(filename)
                extractor = FileExtractor(extract_data.BASE_DIR)
                rows = extractor.read(filename, col_map)
                self.assertGreater(len(rows), 0)
                for raw in rows:
                    out = DataFormatter.format_row(raw, col_map, programas, period)
                    self.assertTrue(out.get("Carrera"))
                    self.assertTrue(out.get("NombreProyecto"))
                    self.assertIsNotNone(out.get("idCodigo"))

    def test_write_output_creates_period_sheet(self):
        filename = self._examples()[0]
        period = ColumnMapper.detect_period(filename)
        col_map = ColumnMapper.get_column_map(filename)
        programas = extract_data.load_programas()
        matcher = FacultadMatcher()
        extractor = FileExtractor(extract_data.BASE_DIR)
        processed = []
        for raw in extractor.read(filename, col_map):
            out = DataFormatter.format_row(raw, col_map, programas, period)
            normalized = matcher.normalize(out.get("Facultad"))
            if normalized:
                out["Facultad"] = normalized
            processed.append(out)
        sheet_name = ColumnMapper.get_sheet_name(period)
        self.assertTrue(extract_data.write_to_output(processed, sheet_name))
        wb = openpyxl.load_workbook(extract_data.OUTPUT_FILE)
        self.assertIn(sheet_name, wb.sheetnames)
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, ColumnMapper.OUTPUT_COLUMNS)
        self.assertGreaterEqual(ws.max_row, 2)
        wb.close()

    def test_cli_end_to_end(self):
        extractor = FileExtractor(extract_data.BASE_DIR)
        filename = self._examples()[0]
        period = ColumnMapper.detect_period(filename)
        sheet_name = ColumnMapper.get_sheet_name(period)
        with mock.patch("extract_data.show_menu", return_value=filename):
            extract_data.main()
        wb = openpyxl.load_workbook(extract_data.OUTPUT_FILE)
        self.assertIn(sheet_name, wb.sheetnames)
        ws = wb[sheet_name]
        self.assertGreaterEqual(ws.max_row, 2)
        wb.close()


if __name__ == "__main__":
    unittest.main()