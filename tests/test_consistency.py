import os
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from modules.date_comparison import consistency
from utils.column_mapper import ColumnMapper

HEADERS = ColumnMapper.OUTPUT_COLUMNS

URL_A = "https://drive.google.com/file/d/Abc123XYZ/view?usp=sharing"
URL_B = "https://drive.google.com/file/d/Def456UVW/view?usp=sharing"
NAME = "Fortalecimiento agrícola"
NAME_INV = "Fortalecimiento Agricola"


class ConsistencyTests(unittest.TestCase):

    def _build_wb(self, sheets):
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in sheets:
            ws = wb.create_sheet(name)
            ws.append(HEADERS)
            for row in rows:
                ws.append([row.get(h) for h in HEADERS])
        return wb

    def _row(self, **overrides):
        row = {
            "NombreProyecto": NAME,
            "idCodigo": "P3AGR18",
            "LinkPlanificacion": URL_A,
            "Carrera": "AGRONOMÍA",
            "Facultad": "CIENCIAS AGRÍCOLAS",
            "Territorio": "Quito",
            "TipoProyecto": "VIGENTE",
        }
        row.update(overrides)
        return row

    def _run(self, wb, extra_fill=None, checks=()):
        if extra_fill:
            for sheet_name, row, color in extra_fill:
                cell = wb[sheet_name].cell(row, 1)
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        with tempfile.TemporaryDirectory() as tmp:
            wb_path = os.path.join(tmp, "BDVinculacionn.xlsx")
            rep_path = os.path.join(tmp, "date_comparison_report.md")
            wb.save(wb_path)
            count = consistency.run(wb_path, rep_path)
            rgb = {}
            for sheet_name, row in checks:
                out = load_workbook(wb_path, data_only=True)
                rgb[(sheet_name, row)] = out[sheet_name].cell(row, 1).fill.start_color.rgb
                out.close()
            report = ""
            if os.path.exists(rep_path):
                with open(rep_path, encoding="utf-8") as f:
                    report = f.read()
            return count, rgb, report

    def test_code_change_across_sheets_reference_is_most_recent(self):
        old = self._row(idCodigo="P3AGR18")
        new = self._row(idCodigo="P3AGR18A")
        wb = self._build_wb([
            ("2025-2025", [old]),
            ("2025-2026", [old]),
            ("2026-2026", [new]),
        ])
        count, rgb, report = self._run(wb, checks=[
            ("2025-2025", 2), ("2025-2026", 2), ("2026-2026", 2),
        ])

        self.assertEqual(count, 2)
        self.assertEqual(rgb[("2025-2025", 2)], "00FFFF00")
        self.assertEqual(rgb[("2025-2026", 2)], "00FFFF00")
        self.assertEqual(rgb[("2026-2026", 2)], "00000000")
        self.assertIn("P3AGR18' differs from 'P3AGR18A'", report)
        self.assertIn("## Consistency across period sheets (yellow)", report)

    def test_older_different_value_flagged_against_most_recent(self):
        same = self._row(idCodigo="P3AGR18", NombreProyecto=NAME_INV)
        latest = self._row(idCodigo="P3AGR18X")
        wb = self._build_wb([
            ("2025-2025", [same]),
            ("2025-2026", [same]),
            ("2026-2026", [latest]),
        ])
        count, rgb, _ = self._run(wb, checks=[
            ("2025-2025", 2), ("2025-2026", 2), ("2026-2026", 2),
        ])
        self.assertEqual(count, 2)
        self.assertEqual(rgb[("2025-2025", 2)], "00FFFF00")
        self.assertEqual(rgb[("2025-2026", 2)], "00FFFF00")
        self.assertEqual(rgb[("2026-2026", 2)], "00000000")

    def test_red_cell_not_overwritten_to_yellow(self):
        old = self._row(idCodigo="P3AGR18")
        new = self._row(idCodigo="P3AGR18B")
        wb = self._build_wb([("2025-2026", [old]), ("2026-2026", [new])])
        count, rgb, _ = self._run(
            wb,
            extra_fill=[("2025-2026", 2, "FF0000")],
            checks=[("2025-2026", 2), ("2026-2026", 2)],
        )
        self.assertEqual(count, 1)
        self.assertEqual(rgb[("2025-2026", 2)], "00FF0000")
        self.assertEqual(rgb[("2026-2026", 2)], "00000000")

    def test_single_occurrence_no_conflict(self):
        wb = self._build_wb([("2026-2026", [self._row()])])
        count, _, _ = self._run(wb)
        self.assertEqual(count, 0)

    def test_link_difference_flagged(self):
        old = self._row(LinkPlanificacion=URL_A)
        latest = self._row(LinkPlanificacion=URL_B)
        wb = self._build_wb([("2025-2026", [old]), ("2026-2026", [latest])])
        count, rgb, _ = self._run(wb, checks=[("2025-2026", 2), ("2026-2026", 2)])
        self.assertEqual(count, 1)
        self.assertEqual(rgb[("2025-2026", 2)], "00FFFF00")
        self.assertEqual(rgb[("2026-2026", 2)], "00000000")


if __name__ == "__main__":
    unittest.main()