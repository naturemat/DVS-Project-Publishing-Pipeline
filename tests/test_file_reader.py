import os
import tempfile
import unittest
from openpyxl import Workbook
from openpyxl.worksheet.hyperlink import Hyperlink

from utils.file_reader import FileExtractor

LINK_COLS = {
    "NombreProyecto": 2,
    "LinkPlanificacion": 19,
}

URL = "https://drive.google.com/file/d/Xyz789ABC/view?usp=sharing"


class FileReaderHyperlinkTests(unittest.TestCase):

    def test_hyperlink_target_read_for_url_cell(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(3, 1, "1")
        cell = ws.cell(3, 19, "Planificación.pdf")
        cell.hyperlink = Hyperlink(ref="S3", target=URL)
        ws.cell(3, 2, "Proyecto ejemplo")

        with tempfile.TemporaryDirectory() as tmp:
            fuente = os.path.join(tmp, "Fuente_Datos")
            os.makedirs(fuente)
            path = os.path.join(fuente, "PERIODO 26-26.xlsx")
            wb.save(path)
            wb.close()

            rows = FileExtractor(tmp).read("PERIODO 26-26.xlsx", LINK_COLS)
            self.assertEqual(rows[0][19], URL)

    def test_non_http_hyperlink_keeps_cell_value(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(3, 1, "1")
        cell = ws.cell(3, 19, "FALTA")
        cell.hyperlink = Hyperlink(ref="S3", target="docs/local/plan.pdf")
        ws.cell(3, 2, "Proyecto ejemplo")

        with tempfile.TemporaryDirectory() as tmp:
            fuente = os.path.join(tmp, "Fuente_Datos")
            os.makedirs(fuente)
            path = os.path.join(fuente, "PERIODO 26-26.xlsx")
            wb.save(path)
            wb.close()

            rows = FileExtractor(tmp).read("PERIODO 26-26.xlsx", LINK_COLS)
            self.assertEqual(rows[0][19], "FALTA")

    def test_plain_url_value_kept(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(3, 1, "1")
        ws.cell(3, 19, URL)
        ws.cell(3, 2, "Proyecto ejemplo")

        with tempfile.TemporaryDirectory() as tmp:
            fuente = os.path.join(tmp, "Fuente_Datos")
            os.makedirs(fuente)
            path = os.path.join(fuente, "PERIODO 26-26.xlsx")
            wb.save(path)
            wb.close()

            rows = FileExtractor(tmp).read("PERIODO 26-26.xlsx", LINK_COLS)
            self.assertEqual(rows[0][19], URL)


if __name__ == "__main__":
    unittest.main()