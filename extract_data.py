import os
import json
import copy
import shutil
import unicodedata
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

from utils.column_mapper import ColumnMapper
from utils.facultad_matcher import FacultadMatcher
from utils.data_formatter import DataFormatter
from utils.file_reader import FileExtractor

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BD_SOURCE = os.path.join(BASE_DIR, "BDVinculacionn 2024-2025.xlsx")
PROGRAMAS_FILE = os.path.join(BASE_DIR, "Programas.json")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "BDVinculacionn.xlsx")

RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
HEADER_FONT = Font(name="Aptos Narrow", size=10, bold=True)
DATA_FONT = Font(name="Aptos Narrow", size=10)
HYPERLINK_FONT = Font(name="Aptos Narrow", size=10, color="0563C1", underline="single")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def load_programas():
    with open(PROGRAMAS_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def _norm_sort_key(value):
    if not value:
        return ""
    s = unicodedata.normalize("NFD", str(value).casefold())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def _proyecto_sort_key(value):
    norm = _norm_sort_key(value)
    if norm and not norm[0].isalnum():
        return "\x00" + norm
    return norm


def show_menu(files):
    print("\n" + "=" * 60)
    print("  EXCEL DATA EXTRACTOR - FUENTE DE DATOS")
    print("=" * 60)
    print("\n  Available files in Fuente_Datos:\n")
    for i, f in enumerate(files, 1):
        print(f"  [{i}] {f}")
    print(f"  [0] Exit\n")
    while True:
        try:
            choice = int(input("  Select a file (number): "))
            if choice == 0:
                return None
            if 1 <= choice <= len(files):
                return files[choice - 1]
            print(f"  Invalid choice. Enter 0-{len(files)}")
        except ValueError:
            print("  Enter a valid number.")


def ensure_output_file():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    if os.path.exists(OUTPUT_FILE):
        return True
    if os.path.exists(BD_SOURCE):
        shutil.copy2(BD_SOURCE, OUTPUT_FILE)
        return True
    return False


def apply_sheet_format(wb, ws, num_cols, num_rows):
    src = None
    for candidate in ["2024-2025", "2023-2024"]:
        if candidate in wb.sheetnames:
            src = wb[candidate]
            break

    if src:
        for col_idx in range(1, num_cols + 1):
            letter = get_column_letter(col_idx)
            if letter in src.column_dimensions:
                ws.column_dimensions[letter].width = src.column_dimensions[letter].width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(num_cols)}1"
    ws.row_dimensions[1].height = 30

    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(1, col_idx)
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        if src:
            cell.border = copy.copy(src.cell(1, col_idx).border)
        else:
            cell.border = THIN_BORDER

    for r in range(2, num_rows + 2):
        ws.row_dimensions[r].height = 27
        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(r, col_idx)
            cell.font = DATA_FONT
            cell.alignment = DATA_ALIGN
            if src:
                cell.border = copy.copy(src.cell(2, col_idx).border)
            else:
                cell.border = THIN_BORDER


def write_to_output(processed, sheet_name):
    if not ensure_output_file():
        print(f"  [ERROR] Could not create output file")
        return False

    wb = openpyxl.load_workbook(OUTPUT_FILE)

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in range(ws.max_row, 0, -1):
            ws.delete_rows(row)
    else:
        ws = wb.create_sheet(sheet_name)

    for col_idx, col_name in enumerate(ColumnMapper.OUTPUT_COLUMNS, 1):
        ws.cell(1, col_idx, col_name)

    sorted_processed = sorted(
        processed,
        key=lambda x: (_norm_sort_key(x.get("Facultad")), _proyecto_sort_key(x.get("NombreProyecto"))),
    )

    for i, row_data in enumerate(sorted_processed):
        r = i + 2
        row_data["idDocumento"] = i + 1
        for col_idx, col_name in enumerate(ColumnMapper.OUTPUT_COLUMNS, 1):
            ws.cell(r, col_idx, row_data.get(col_name))

        has_missing = len(row_data.get("_missing", [])) > 0
        program_found = row_data.get("_program_info_found", False)

        if has_missing:
            for col_idx in range(1, len(ColumnMapper.OUTPUT_COLUMNS) + 1):
                ws.cell(r, col_idx).fill = RED_FILL
        elif program_found:
            for col_idx in range(1, len(ColumnMapper.OUTPUT_COLUMNS) + 1):
                ws.cell(r, col_idx).fill = ORANGE_FILL

    apply_sheet_format(wb, ws, len(ColumnMapper.OUTPUT_COLUMNS), len(sorted_processed))

    link_planif_col = ColumnMapper.OUTPUT_COLUMNS.index("LinkPlanificacion") + 1
    for r in range(2, len(sorted_processed) + 2):
        cell = ws.cell(r, link_planif_col)
        url = cell.value
        if url and str(url).startswith("http"):
            cell.hyperlink = str(url)
            cell.font = HYPERLINK_FONT

    wb.save(OUTPUT_FILE)
    wb.close()
    return True


def generate_report(missing_report, filename, sheet_name):
    if not missing_report:
        return

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    report_path = os.path.join(OUTPUT_DIR, "missing_data_report.md")

    lines = [
        "# Missing Data Report",
        "",
        f"**Source File:** {filename}",
        f"**Target Sheet:** {sheet_name}",
        f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "## Rows with Missing Data",
        "",
        "| Source Row | Codigo | Nombre Proyecto | Missing Fields |",
        "|------------|--------|-----------------|----------------|",
    ]

    for item in missing_report:
        nombre = str(item.get("nombre", "N/A"))[:50]
        missing = ", ".join(item["missing"])
        lines.append(
            f"| {item['source_row']} | {item.get('codigo', 'N/A')} | {nombre} | {missing} |"
        )

    lines.append("")
    lines.append(f"**Total rows with missing data:** {len(missing_report)}")

    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print(f"  Report saved: {report_path}")


def main():
    extractor = FileExtractor(BASE_DIR)
    files = extractor.list_files()

    if not files:
        print("\n  No Excel files found in Fuente_Datos folder.")
        return

    filename = show_menu(files)
    if not filename:
        print("\n  Goodbye!")
        return

    print(f"\n  Processing: {filename}")

    period = ColumnMapper.detect_period(filename)
    if not period:
        print("  [ERROR] Cannot determine period from filename.")
        return

    sheet_name = ColumnMapper.get_sheet_name(period)
    print(f"  Target sheet: {sheet_name}")

    col_map = ColumnMapper.get_column_map(filename)
    if not col_map:
        print(f"  [ERROR] No column mapping for {filename}")
        return

    print("  Loading Programas.json...")
    programas = load_programas()

    print("  Loading facultades.json...")
    facultad_matcher = FacultadMatcher()

    print("  Reading source data...")
    raw_rows = extractor.read(filename, col_map)
    print(f"  Found {len(raw_rows)} data rows.")

    if not raw_rows:
        print("  No data found. Exiting.")
        return

    print("  Processing data...")
    processed = []
    missing_report = []

    for raw_row in raw_rows:
        out = DataFormatter.format_row(raw_row, col_map, programas, period)

        if out.get("Facultad"):
            normalized = facultad_matcher.normalize(out["Facultad"])
            if normalized:
                out["Facultad"] = normalized

        override = DataFormatter.CARRERA_FACULTAD_OVERRIDE.get(out.get("Carrera"))
        if override:
            out["Facultad"] = override
            if "Facultad" in out.get("_missing", []):
                out["_missing"].remove("Facultad")

        out["_source_row"] = raw_row.get("_row_idx")
        out["idCodigo"] = out.get("idCodigo")

        if out["_missing"]:
            missing_report.append({
                "source_row": raw_row.get("_row_idx"),
                "codigo": out.get("idCodigo"),
                "nombre": out.get("NombreProyecto"),
                "missing": out["_missing"]
            })

        processed.append(out)

    print(f"  Processed {len(processed)} rows.")
    print(f"  Rows with missing data: {len(missing_report)}")

    rows_with_program_lookup = sum(1 for p in processed if p.get("_program_info_found"))
    if rows_with_program_lookup:
        print(f"  Rows enriched from Programas.json: {rows_with_program_lookup}")

    print(f"  Writing to sheet '{sheet_name}'...")
    success = write_to_output(processed, sheet_name)

    if success:
        print(f"  Data written to: {OUTPUT_FILE}")
    else:
        print("  Failed to write data.")
        return

    if missing_report:
        print("  Generating missing data report...")
        generate_report(missing_report, filename, sheet_name)

    print("\n" + "=" * 60)
    print("  SUMMARY")
    print("=" * 60)
    print(f"  Source: {filename}")
    print(f"  Sheet: {sheet_name}")
    print(f"  Output: {OUTPUT_FILE}")
    print(f"  Total rows: {len(processed)}")
    print(f"  Complete rows: {len(processed) - len(missing_report)}")
    print(f"  Missing data rows: {len(missing_report)} (marked RED)")
    print(f"  Programas.json lookups: {rows_with_program_lookup} (marked ORANGE)")
    print("=" * 60)


if __name__ == "__main__":
    main()
