import openpyxl
import os


class FileExtractor:

    def __init__(self, base_dir):
        self.base_dir = base_dir
        self.fuente_dir = os.path.join(base_dir, "Fuente_Datos")

    def list_files(self):
        files = [f for f in os.listdir(self.fuente_dir)
                 if f.endswith(".xlsx") and not f.startswith("~$")]
        return sorted(files)

    def read(self, filename, col_map):
        filepath = os.path.join(self.fuente_dir, filename)
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active

        data_start = 3
        rows = []
        for row_idx in range(data_start, ws.max_row + 1):
            first_val = ws.cell(row_idx, 1).value
            if first_val is None:
                continue
            try:
                if float(first_val) == 0:
                    continue
            except (ValueError, TypeError):
                pass

            row_data = {}
            all_cols = set(col_map.values()) - {None}
            max_col = max(all_cols) if all_cols else 0
            for col in range(1, max_col + 1):
                cell = ws.cell(row_idx, col)
                value = cell.value
                if cell.hyperlink and isinstance(cell.hyperlink.target, str):
                    if cell.hyperlink.target.startswith("http"):
                        value = cell.hyperlink.target
                row_data[col] = value
            row_data["_row_idx"] = row_idx
            rows.append(row_data)

        wb.close()
        return rows
