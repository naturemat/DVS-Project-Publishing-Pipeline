import re
import unicodedata
from openpyxl.styles import PatternFill

RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

RED_RGB = ("FFFF0000", "00FF0000")

REQUIRED_FIELDS = (
    "Facultad", "Carrera", "NombreProyecto", "NombrePrograma",
    "NombreCoordinador", "Territorio", "FechaInicio", "FechaFin", "idCodigo",
)

_FULL_DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")


def normalized(value):
    if not value:
        return ""
    s = unicodedata.normalize("NFD", str(value).casefold())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join(s.split())


def is_red(cell):
    rgb = getattr(cell.fill.start_color, "rgb", None)
    return rgb in RED_RGB


def missing_required(ws, r, headers):
    missing = []
    for field in REQUIRED_FIELDS:
        col = headers.get(field)
        if not col:
            missing.append(field)
            continue
        value = ws.cell(r, col).value
        if value is None or not str(value).strip():
            missing.append(field)
        elif field in ("FechaInicio", "FechaFin") and not _FULL_DATE_RE.match(str(value).strip()):
            missing.append(field)
    return missing


def mark_corrected(ws, r, num_cols, headers):
    fill = RED_FILL if missing_required(ws, r, headers) else ORANGE_FILL
    for c in range(1, num_cols + 1):
        ws.cell(r, c).fill = fill


def mark_yellow(ws, r, num_cols):
    for c in range(1, num_cols + 1):
        if not is_red(ws.cell(r, c)):
            ws.cell(r, c).fill = YELLOW_FILL