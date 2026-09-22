import os
import re
import unicodedata
from datetime import datetime, date
from openpyxl import load_workbook
from openpyxl.styles import Font

from utils.column_mapper import ColumnMapper

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "BDVinculacionn.xlsx")
REPORT_FILE = os.path.join(OUTPUT_DIR, "date_comparison_report.md")

FILLABLE_FIELDS = (
    "idCodigo", "LinkPlanificacion", "Facultad", "Carrera",
    "NombrePrograma", "NombreCoordinador", "Territorio",
    "FechaInicio", "FechaFin",
)
FILL_TARGET_TYPES = ("VIGENTE",)
_FULL_DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
EMPTY_TEXT = {"N/A", "NA"}
HYPERLINK_FONT = Font(name="Aptos Narrow", size=10, color="0563C1", underline="single")


def _headers(ws):
    return {
        cell.value: idx
        for idx, cell in enumerate(ws[1], 1)
        if cell.value
    }


def _norm_name(value):
    if not value:
        return None
    s = unicodedata.normalize("NFD", str(value).casefold())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join(s.split())


def _is_empty(value):
    if value is None:
        return True
    s = str(value).strip()
    return not s or s.upper() in EMPTY_TEXT


def _clean_source(field, value):
    if _is_empty(value):
        return None
    if field in ("FechaInicio", "FechaFin"):
        if isinstance(value, (datetime, date)):
            return value.strftime("%d/%m/%Y")
        s = str(value).strip()
        if not _FULL_DATE_RE.match(s):
            return None
        return s
    if isinstance(value, str):
        return value.strip()
    return value


def _collect_sheet(ws, headers):
    projects = {}
    fillable = [f for f in FILLABLE_FIELDS if headers.get(f)]
    name_col = headers.get("NombreProyecto")
    if not name_col:
        return projects
    for r in range(2, ws.max_row + 1):
        norm = _norm_name(ws.cell(r, name_col).value)
        if not norm:
            continue
        entry = projects.setdefault(norm, {"_rows": []})
        entry["_rows"].append((ws.title, r))
        for field in fillable:
            if field in entry:
                continue
            value = _clean_source(field, ws.cell(r, headers[field]).value)
            if value is not None:
                entry[field] = value
    return projects


def _fill_sheet(ws, headers, index, changes):
    fillable = [f for f in FILLABLE_FIELDS if headers.get(f)]
    name_col = headers.get("NombreProyecto")
    tipo_col = headers.get("TipoProyecto")
    if not name_col or not fillable:
        return 0
    cells_filled = 0
    for r in range(2, ws.max_row + 1):
        nombre = ws.cell(r, name_col).value
        norm = _norm_name(nombre)
        if not norm:
            continue
        target_type = str(ws.cell(r, tipo_col).value or "").strip().upper() if tipo_col else ""
        if target_type not in FILL_TARGET_TYPES:
            continue
        entry = index.get(norm)
        if not entry:
            continue
        filled = []
        for field in fillable:
            cell = ws.cell(r, headers[field])
            if not _is_empty(cell.value):
                continue
            value = entry.get(field)
            if value is None:
                continue
            if isinstance(value, str):
                value = value.strip()
            cell.value = value
            if field == "LinkPlanificacion" and str(value).startswith("http"):
                cell.hyperlink = str(value)
                cell.font = HYPERLINK_FONT
            filled.append(field)
        if filled:
            cells_filled += len(filled)
            changes.append({
                "sheet": ws.title,
                "row": r,
                "nombre": str(nombre or "")[:60],
                "fields": filled,
            })
    return cells_filled


def run(workbook_path=None, report_path=None):
    wb_path = workbook_path or OUTPUT_FILE
    rep_path = report_path or REPORT_FILE

    if not os.path.exists(wb_path):
        print("  AUTOFILL: output workbook not found, skipping.", flush=True)
        return 0

    wb = load_workbook(wb_path)
    managed = set(ColumnMapper.PERIOD_SHEET_MAP.values())
    available = [s for s in wb.sheetnames if s in managed]

    index = {}
    for sheet_name in available:
        ws = wb[sheet_name]
        headers = _headers(ws)
        for norm, entry in _collect_sheet(ws, headers).items():
            base = index.setdefault(norm, {"_rows": []})
            for key, value in entry.items():
                if key == "_rows":
                    base["_rows"].extend(value)
                elif key not in base:
                    base[key] = value

    changes = []
    total_filled = 0
    for sheet_name in available:
        ws = wb[sheet_name]
        total_filled += _fill_sheet(ws, _headers(ws), index, changes)

    if changes:
        wb.save(wb_path)
    wb.close()

    _append_report(rep_path, changes, total_filled)
    _print_summary(changes, total_filled, len(available))
    return total_filled


def _print_summary(changes, total_filled, sheet_count):
    print("\n  AUTOFILL: complete missing fields from other sheets (VIGENTE projects only)", flush=True)
    if not changes:
        print("  No VIGENTE row had fillable missing fields.", flush=True)
        return
    per_sheet = {}
    for item in changes:
        per_sheet.setdefault(item["sheet"], []).append(item)
    for sheet_name, items in per_sheet.items():
        print(f"    {sheet_name}: {len(items)} row(s) completed", flush=True)
    print(f"  Total cells filled: {total_filled} across {len(changes)} row(s) in {sheet_count} period sheet(s)", flush=True)


def _append_report(rep_path, changes, total_filled):
    if not changes:
        return
    lines = ["", "## Autofill (same NombreProyecto across sheets, VIGENTE only)", ""]
    for item in changes:
        lines.append(
            f"- **{item['sheet']}** row {item['row']} - {item['nombre']}: "
            f"filled {', '.join(item['fields'])}"
        )
    lines.append("")
    lines.append(f"**Total cells filled:** {total_filled}")
    with open(rep_path, "a", encoding="utf-8") as f:
        f.write("\n".join(lines))